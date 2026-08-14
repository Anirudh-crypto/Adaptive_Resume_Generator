from __future__ import annotations

import datetime
import io

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from app.models import Contact, CoverLetter, Resume

DOCX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

BODY_FONT = "Calibri"
BODY_SIZE = Pt(11)
NAME_SIZE = Pt(20)
CONTACT_SIZE = Pt(9.5)
MUTED = RGBColor(0x44, 0x44, 0x44)


def _contact_line(contact: Contact) -> str:
    parts = [
        contact.location,
        contact.phone,
        contact.email,
        contact.linkedin,
        contact.github,
        contact.website,
    ]
    return "  |  ".join(p for p in parts if p)


def _paragraph(document: DocxDocument, text: str = "", *, space_after: float = 10.0):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(space_after)
    if text:
        paragraph.add_run(text)
    return paragraph


def _today() -> str:
    # %-d / %#d are platform-specific, so strip the zero padding by hand.
    today = datetime.date.today()
    return f"{today.strftime('%B')} {today.day}, {today.year}"


def render_cover_letter_docx(letter: CoverLetter, resume: Resume) -> bytes:
    """Render a cover letter to .docx bytes. Identity (name, contact details) comes from the
    resume, never from the model output — same anti-hallucination split as the resume pipeline."""
    document = Document()

    normal = document.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = BODY_SIZE
    normal.paragraph_format.space_after = Pt(10)
    normal.paragraph_format.line_spacing = 1.15

    for section in document.sections:
        section.top_margin = section.bottom_margin = Pt(72)
        section.left_margin = section.right_margin = Pt(72)

    name_paragraph = _paragraph(document, space_after=2.0)
    name_run = name_paragraph.add_run(resume.name)
    name_run.bold = True
    name_run.font.size = NAME_SIZE

    contact_line = _contact_line(resume.contact)
    if contact_line:
        contact_paragraph = _paragraph(document, space_after=18.0)
        contact_run = contact_paragraph.add_run(contact_line)
        contact_run.font.size = CONTACT_SIZE
        contact_run.font.color.rgb = MUTED

    _paragraph(document, _today(), space_after=18.0)

    addressee = [line for line in (letter.company, _role_line(letter)) if line]
    if addressee:
        for line in addressee[:-1]:
            _paragraph(document, line, space_after=0.0)
        _paragraph(document, addressee[-1], space_after=18.0)

    _paragraph(document, letter.salutation, space_after=12.0)

    for body_paragraph in letter.paragraphs:
        paragraph = _paragraph(document, body_paragraph, space_after=12.0)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

    _paragraph(document, letter.closing, space_after=24.0)
    _paragraph(document, resume.name, space_after=0.0)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _role_line(letter: CoverLetter) -> str:
    return f"Re: {letter.role}" if letter.role else ""
