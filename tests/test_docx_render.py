from __future__ import annotations

import io
from pathlib import Path

from docx import Document

from app.docx_render import render_cover_letter_docx
from app.models import CoverLetter
from app.resume_parser import parse_resume

FIXTURE = Path(__file__).parent / "fixtures" / "sample_resume.md"

LETTER = CoverLetter(
    company="Globex Corporation",
    role="Staff Backend Engineer",
    salutation="Dear Hiring Manager,",
    paragraphs=[
        "I am applying for the Staff Backend Engineer role at Globex Corporation.",
        "At Acme Corp I led a request-routing service that improved p99 latency by 30%.",
        "I would welcome the chance to talk about the role.",
    ],
    closing="Sincerely,",
)


def _texts(docx_bytes: bytes) -> list[str]:
    document = Document(io.BytesIO(docx_bytes))
    return [p.text for p in document.paragraphs]


def test_render_produces_a_readable_docx():
    docx_bytes = render_cover_letter_docx(LETTER, parse_resume(FIXTURE))
    assert docx_bytes[:2] == b"PK"  # .docx is a zip container

    texts = _texts(docx_bytes)
    for paragraph in LETTER.paragraphs:
        assert paragraph in texts
    assert "Dear Hiring Manager," in texts
    assert "Sincerely," in texts
    assert "Globex Corporation" in texts
    assert "Re: Staff Backend Engineer" in texts


def test_identity_comes_from_the_resume_not_the_letter():
    resume = parse_resume(FIXTURE)
    texts = _texts(render_cover_letter_docx(LETTER, resume))

    assert texts[0] == "Jane Doe"  # name heads the letter
    assert texts[-1] == "Jane Doe"  # and signs it
    contact_line = texts[1]
    assert "jane.doe@example.com" in contact_line
    assert "Seattle, WA" in contact_line
    assert "(555) 123-4567" in contact_line


def test_missing_company_and_role_are_omitted_not_placeholdered():
    letter = LETTER.model_copy(update={"company": "", "role": ""})
    texts = _texts(render_cover_letter_docx(letter, parse_resume(FIXTURE)))

    assert not any("Re:" in t for t in texts)
    assert not any("[" in t for t in texts)
